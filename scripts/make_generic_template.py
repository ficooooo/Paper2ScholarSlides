#!/usr/bin/env python
"""Create a public-safe academic PPT template with no institutional branding."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "assets" / "templates" / "academic-template.pptx"

NAVY = RGBColor(0x0B, 0x3A, 0x5B)
BLUE = RGBColor(0x1D, 0x70, 0xB7)
PALE = RGBColor(0xF4, 0xF8, 0xFB)
TEXT = RGBColor(0x20, 0x2A, 0x33)
MUTED = RGBColor(0x5D, 0x6B, 0x78)


def add_text(slide, text, x, y, w, h, size=20, bold=False, color=TEXT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_header(slide, title, section, page):
    add_text(slide, "Academic Review Decksmith", 0.45, 0.20, 4.8, 0.25, size=10, bold=True, color=MUTED)
    add_text(slide, section, 10.2, 0.20, 2.5, 0.25, size=9, color=MUTED)
    add_text(slide, title, 0.45, 0.62, 9.6, 0.45, size=25, bold=True, color=NAVY)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(1.17), Inches(12.0), Inches(0.03))
    line.fill.solid()
    line.fill.fore_color.rgb = BLUE
    line.line.fill.background()
    footer = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.12), Inches(13.333), Inches(0.38))
    footer.fill.solid()
    footer.fill.fore_color.rgb = NAVY
    footer.line.fill.background()
    add_text(slide, "Academic review presentation | source-aware figures, formulas, and QA", 0.45, 7.22, 8.5, 0.16, size=7.2, color=RGBColor(255, 255, 255))
    add_text(slide, f"{page:02d}", 12.35, 7.20, 0.45, 0.18, size=8, bold=True, color=RGBColor(255, 255, 255))


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    cover = prs.slides.add_slide(blank)
    add_text(cover, "Academic Review Title", 0.75, 1.65, 6.8, 0.55, size=30, bold=True, color=NAVY)
    add_text(cover, "Mechanism, Structure, Modeling, Control, and Applications", 0.75, 2.35, 8.2, 0.38, size=15, color=TEXT)
    cover.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.4), Inches(1.45), Inches(3.8), Inches(2.2)).fill.solid()
    cover.shapes[-1].fill.fore_color.rgb = PALE
    add_text(cover, "Figure Placeholder", 9.25, 2.38, 2.0, 0.25, size=13, color=MUTED)
    add_text(cover, "Presenter / Affiliation / Date", 0.75, 5.55, 5.0, 0.3, size=11, color=MUTED)

    content = prs.slides.add_slide(blank)
    add_header(content, "Slide Title: One Technical Claim", "Section Name", 2)
    add_text(content, "Core claim sentence goes here.", 0.65, 1.42, 11.8, 0.32, size=12, bold=True, color=RGBColor(255, 255, 255))
    bar = content.shapes[-1]
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    add_text(content, "• Evidence point tied to a source, variable, method, or metric.\n• Avoid decorative terms; explain what the figure or formula means.\n• Keep source labels traceable.", 0.78, 2.05, 4.8, 1.25, size=13, color=TEXT)
    fig = content.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.25), Inches(1.95), Inches(5.4), Inches(2.7))
    fig.fill.solid()
    fig.fill.fore_color.rgb = PALE
    fig.line.color.rgb = RGBColor(0xC9, 0xD8, 0xE6)
    add_text(content, "Image / diagram / table", 7.55, 3.05, 2.8, 0.35, size=15, bold=True, color=MUTED)
    note = content.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.75), Inches(5.75), Inches(11.6), Inches(0.58))
    note.fill.solid()
    note.fill.fore_color.rgb = PALE
    note.line.color.rgb = BLUE
    add_text(content, "Read the figure: explain what the audience should compare and why it supports the claim.", 0.93, 5.91, 10.8, 0.22, size=9, color=TEXT)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
