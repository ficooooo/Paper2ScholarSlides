#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Build a refined project poster and clean PPT preview images."""

from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageColor, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
GALLERY = ROOT / "assets" / "gallery"
FIGURES = ROOT / "assets" / "figures"
POSTERS = ROOT / "assets" / "posters"
POSTERS.mkdir(parents=True, exist_ok=True)
GALLERY.mkdir(parents=True, exist_ok=True)

REPO_NAME = "Paper2ScholarSlides"
REPO_URL = "https://github.com/ficooooo/Paper2ScholarSlides"
INTRO = (
    "面向学术综述与研究汇报的幻灯片构建技能。"
    "它将论文资料、综述初稿与既有模板整理为结构清晰、图表来源明确、"
    "公式可解释、版式可复核的研究型 PPT 资产，适用于课程汇报、组会交流、"
    "开题答辩与专题综述展示。"
)

FEATURES = [
    ("结构不是堆章节", "按“背景—问题—方法—建模—控制—应用”重组内容，让综述页序真正服务于讲述逻辑。", "#2B6CB0"),
    ("图片必须可追溯", "区分文献原图、重绘示意、原生表格与展示性配图，避免整套幻灯片视觉来源混乱。", "#1F8A70"),
    ("公式需要可读性", "关键公式不仅展示，还补充变量、物理含义、适用边界和与图表的对应关系。", "#D97706"),
    ("交付前做导出核查", "统一导出 PNG 逐页检查溢出、遮挡、低清晰度、占位文本和引用说明缺失问题。", "#D64550"),
]

WORKFLOW = ["资料输入", "论证重构", "图表处理", "PPT 成稿", "导出核验"]
WORKFLOW_COLORS = ["#2B6CB0", "#1F8A70", "#D97706", "#E24A59", "#143C66"]

USE_CASES = ["课程综述汇报", "课题组组会", "博士生预答辩", "文献综述展示"]
ASSET_LINES = [
    "PPT 预览图：真实页面重排导出",
    "海报 PNG：适合 GitHub 展示与汇报封面",
    "海报 PPT：可继续替换文案与图片",
    "图件资产：控制框图、线图与结构示意",
]

PREVIEW_SLIDES = [
    (3, "研究背景"),
    (8, "典型步态"),
    (11, "关节布置"),
    (14, "运动建模"),
    (18, "控制总览"),
    (23, "管道应用"),
]

DETAIL_SLIDES = [
    (11, "结构页"),
    (14, "建模页"),
    (20, "控制页"),
]


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    path = "C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc"
    return ImageFont.truetype(path, size)


def rgb(color: str) -> tuple[int, int, int]:
    return ImageColor.getrgb(color)


def rgba(color: str, alpha: int) -> tuple[int, int, int, int]:
    r, g, b = rgb(color)
    return (r, g, b, alpha)


def vertical_gradient(size: tuple[int, int], top: str, bottom: str) -> Image.Image:
    width, height = size
    img = Image.new("RGBA", size)
    top_rgb = rgb(top)
    bottom_rgb = rgb(bottom)
    pixels = img.load()
    for y in range(height):
        ratio = y / max(1, height - 1)
        c = tuple(int(top_rgb[i] * (1 - ratio) + bottom_rgb[i] * ratio) for i in range(3)) + (255,)
        for x in range(width):
            pixels[x, y] = c
    return img


def rounded_panel(
    base: Image.Image,
    box: tuple[int, int, int, int],
    fill: str = "#FFFFFF",
    outline: str = "#D8E4F0",
    radius: int = 28,
    shadow: bool = True,
) -> None:
    x1, y1, x2, y2 = box
    if shadow:
        shadow_img = Image.new("RGBA", base.size, (0, 0, 0, 0))
        shadow_draw = ImageDraw.Draw(shadow_img)
        shadow_draw.rounded_rectangle(
            (x1 + 8, y1 + 12, x2 + 8, y2 + 12),
            radius=radius,
            fill=(8, 24, 48, 24),
        )
        shadow_img = shadow_img.filter(ImageFilter.GaussianBlur(14))
        base.alpha_composite(shadow_img)
    draw = ImageDraw.Draw(base)
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=2)


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont, max_width: int) -> list[str]:
    lines: list[str] = []
    for paragraph in text.split("\n"):
        current = ""
        for ch in paragraph:
            trial = current + ch
            width = draw.textbbox((0, 0), trial, font=font)[2]
            if width <= max_width or not current:
                current = trial
            else:
                lines.append(current)
                current = ch
        if current:
            lines.append(current)
    return lines


def draw_wrapped(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int],
    text: str,
    font: ImageFont.FreeTypeFont,
    fill: str,
    max_width: int,
    line_gap: int = 8,
) -> int:
    x, y = xy
    h = 0
    for line in wrap_text(draw, text, font, max_width):
        draw.text((x, y + h), line, font=font, fill=fill)
        bbox = draw.textbbox((0, 0), line if line else "A", font=font)
        h += (bbox[3] - bbox[1]) + line_gap
    return h


def fit_image_cover(img: Image.Image, target: tuple[int, int]) -> Image.Image:
    target_w, target_h = target
    src_ratio = img.width / img.height
    dst_ratio = target_w / target_h
    if src_ratio > dst_ratio:
        new_h = target_h
        new_w = int(target_h * src_ratio)
    else:
        new_w = target_w
        new_h = int(target_w / src_ratio)
    resized = img.resize((new_w, new_h), Image.Resampling.LANCZOS)
    left = (new_w - target_w) // 2
    top = (new_h - target_h) // 2
    return resized.crop((left, top, left + target_w, top + target_h))


def fit_image_contain(
    img: Image.Image,
    target: tuple[int, int],
    background: str = "#FFFFFF",
    padding: int = 0,
) -> Image.Image:
    target_w, target_h = target
    inner_w = max(1, target_w - padding * 2)
    inner_h = max(1, target_h - padding * 2)
    src_ratio = img.width / img.height
    dst_ratio = inner_w / inner_h
    if src_ratio > dst_ratio:
        new_w = inner_w
        new_h = int(inner_w / src_ratio)
    else:
        new_h = inner_h
        new_w = int(inner_h * src_ratio)
    resized = img.resize((new_w, new_h), Image.Resampling.LANCZOS)
    canvas = Image.new("RGB", (target_w, target_h), background)
    x = (target_w - new_w) // 2
    y = (target_h - new_h) // 2
    canvas.paste(resized, (x, y))
    return canvas


def source_contact_sheet() -> Image.Image:
    return Image.open(GALLERY / "snake-robot-demo-overview-real.png").convert("RGB")


def slide_crop(sheet: Image.Image, slide_number: int) -> Image.Image:
    cols, rows = 3, 9
    cell_w = sheet.width // cols
    cell_h = sheet.height // rows
    index = slide_number - 1
    row = index // cols
    col = index % cols
    x1 = col * cell_w
    y1 = row * cell_h
    crop = sheet.crop((x1 + 28, y1 + 72, x1 + cell_w - 28, y1 + cell_h - 26))
    return crop


def clean_preview_grid() -> Path:
    sheet = source_contact_sheet()
    canvas = vertical_gradient((1920, 900), "#F7FAFD", "#EEF5FB")
    overlay = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
    overlay_draw = ImageDraw.Draw(overlay)
    overlay_draw.ellipse((1280, -120, 1980, 480), fill=rgba("#BEDDFF", 48))
    overlay_draw.ellipse((-120, 620, 520, 1120), fill=rgba("#DCEAFE", 58))
    overlay = overlay.filter(ImageFilter.GaussianBlur(18))
    canvas.alpha_composite(overlay)
    draw = ImageDraw.Draw(canvas)
    draw.text((72, 52), "PPT Preview", font=load_font(54, bold=True), fill="#103A5D")
    draw.text((72, 120), "基于真实汇报页面重排生成的精选预览图", font=load_font(26), fill="#5A7086")

    panel = (58, 170, 1862, 842)
    rounded_panel(canvas, panel, fill="#FFFFFF", outline="#D8E4F0", radius=34, shadow=True)
    start_x = 96
    start_y = 220
    tile_w, tile_h = 550, 312
    x_gap, y_gap = 26, 24

    for idx, (slide_no, label) in enumerate(PREVIEW_SLIDES):
        row = idx // 3
        col = idx % 3
        x = start_x + col * (tile_w + x_gap)
        y = start_y + row * (tile_h + y_gap)
        card = (x, y, x + tile_w, y + tile_h)
        rounded_panel(canvas, card, fill="#FBFDFF", outline="#D6E1EC", radius=24, shadow=False)
        preview = fit_image_contain(slide_crop(sheet, slide_no), (tile_w - 24, 220), background="#FFFFFF", padding=8)
        mask = Image.new("L", (tile_w - 24, 220), 0)
        ImageDraw.Draw(mask).rounded_rectangle((0, 0, tile_w - 25, 219), radius=18, fill=255)
        canvas.paste(preview, (x + 12, y + 12), mask)
        draw.text((x + 16, y + 244), f"Slide {slide_no}", font=load_font(18), fill="#8093A8")
        draw.text((x + 16, y + 268), label, font=load_font(24, bold=True), fill="#133E66")

    out = GALLERY / "paper2scholarslides-ppt-preview-clean.png"
    canvas.convert("RGB").save(out, quality=96)
    return out


def detail_strip() -> Path:
    sheet = source_contact_sheet()
    canvas = Image.new("RGBA", (1250, 350), "#F5F8FC")
    cards = []
    for slide_no, label in DETAIL_SLIDES:
        card = Image.new("RGBA", (390, 306), (255, 255, 255, 255))
        cdraw = ImageDraw.Draw(card)
        cdraw.rounded_rectangle((0, 0, 389, 305), radius=28, fill="#FFFFFF", outline="#D6E1EC", width=2)
        preview = fit_image_contain(slide_crop(sheet, slide_no), (364, 206), background="#FFFFFF", padding=8)
        mask = Image.new("L", (364, 206), 0)
        ImageDraw.Draw(mask).rounded_rectangle((0, 0, 363, 205), radius=18, fill=255)
        card.paste(preview, (13, 16), mask)
        cdraw.text((18, 238), f"Slide {slide_no}", font=load_font(16), fill="#7B8FA4")
        cdraw.text((18, 258), label, font=load_font(22, bold=True), fill="#133E66")
        cards.append(card)
    for i, card in enumerate(cards):
        canvas.alpha_composite(card, (i * 416, 22))
    out = GALLERY / "paper2scholarslides-ppt-preview-strip.png"
    canvas.convert("RGB").save(out, quality=96)
    return out


def feature_card(base: Image.Image, box: tuple[int, int, int, int], title: str, body: str, accent: str) -> None:
    rounded_panel(base, box, fill="#FFFFFF", outline="#D8E4F0", radius=28, shadow=True)
    draw = ImageDraw.Draw(base)
    x1, y1, x2, y2 = box
    draw.rounded_rectangle((x1 + 18, y1 + 18, x1 + 48, y1 + 48), radius=8, fill=accent)
    draw.text((x1 + 62, y1 + 14), title, font=load_font(22, bold=True), fill="#123B61")
    draw_wrapped(draw, (x1 + 18, y1 + 62), body, load_font(18), "#314252", x2 - x1 - 36, 6)


def preview_panel(base: Image.Image, box: tuple[int, int, int, int], title: str, subtitle: str, img: Image.Image) -> None:
    rounded_panel(base, box, fill="#FFFFFF", outline="#D8E4F0", radius=34, shadow=True)
    draw = ImageDraw.Draw(base)
    x1, y1, x2, y2 = box
    draw.text((x1 + 26, y1 + 20), title, font=load_font(32, bold=True), fill="#123B61")
    draw.text((x1 + 26, y1 + 62), subtitle, font=load_font(18), fill="#60768C")
    inner = (x1 + 22, y1 + 104, x2 - 22, y2 - 22)
    fit = fit_image_contain(img, (inner[2] - inner[0], inner[3] - inner[1]), background="#FFFFFF", padding=10)
    mask = Image.new("L", (inner[2] - inner[0], inner[3] - inner[1]), 0)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, mask.width - 1, mask.height - 1), radius=26, fill=255)
    base.paste(fit, (inner[0], inner[1]), mask)


def flow_panel(base: Image.Image, box: tuple[int, int, int, int]) -> None:
    rounded_panel(base, box, fill="#FFFFFF", outline="#D8E4F0", radius=28, shadow=True)
    draw = ImageDraw.Draw(base)
    x1, y1, x2, y2 = box
    draw.text((x1 + 22, y1 + 18), "工作流程", font=load_font(28, bold=True), fill="#123B61")
    total_w = x2 - x1 - 44
    step_w = 112
    gap = (total_w - step_w * 5) // 4
    y = y1 + 106
    for idx, label in enumerate(WORKFLOW):
        bx1 = x1 + 22 + idx * (step_w + gap)
        bx2 = bx1 + step_w
        color = WORKFLOW_COLORS[idx]
        draw.rounded_rectangle((bx1, y, bx2, y + 62), radius=18, fill="#FBFDFF", outline=color, width=2)
        tw = draw.textbbox((0, 0), label, font=load_font(20, bold=True))[2]
        draw.text((bx1 + (step_w - tw) / 2, y + 16), label, font=load_font(20, bold=True), fill=color)
        if idx < 4:
            ax1 = bx2 + 10
            ax2 = bx2 + gap - 12
            ay = y + 31
            draw.line((ax1, ay, ax2, ay), fill="#7CBAD2", width=4)
            draw.polygon([(ax2, ay), (ax2 - 12, ay - 7), (ax2 - 12, ay + 7)], fill="#7CBAD2")
    draw.text((x1 + 24, y1 + 194), "从论文资料到成品海报，强调页面叙事、图表说明与导出核验。", font=load_font(18), fill="#60768C")


def resource_panel(base: Image.Image, box: tuple[int, int, int, int], figure_path: Path) -> None:
    rounded_panel(base, box, fill="#FFFFFF", outline="#D8E4F0", radius=28, shadow=True)
    draw = ImageDraw.Draw(base)
    x1, y1, x2, y2 = box
    draw.text((x1 + 22, y1 + 18), "图件资源与输出", font=load_font(28, bold=True), fill="#123B61")
    figure = Image.open(figure_path).convert("RGB")
    fit = fit_image_contain(figure, (x2 - x1 - 44, 138), background="#FFFFFF", padding=8)
    mask = Image.new("L", (x2 - x1 - 44, 138), 0)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, mask.width - 1, mask.height - 1), radius=18, fill=255)
    base.paste(fit, (x1 + 22, y1 + 64), mask)
    draw.line((x1 + 22, y1 + 220, x2 - 22, y1 + 220), fill="#E2EBF4", width=2)
    for idx, line in enumerate(ASSET_LINES):
        yy = y1 + 238 + idx * 28
        draw.rounded_rectangle((x1 + 22, yy + 7, x1 + 36, yy + 21), radius=4, fill="#2B6CB0")
        draw.text((x1 + 48, yy), line, font=load_font(18), fill="#314252")


def poster_png(preview_path: Path, strip_path: Path) -> Path:
    size = (2400, 1350)
    base = vertical_gradient(size, "#F6FAFD", "#EDF5FB")
    overlay = Image.new("RGBA", size, (0, 0, 0, 0))
    od = ImageDraw.Draw(overlay)
    od.ellipse((1720, -140, 2620, 620), fill=rgba("#CFE3FF", 52))
    od.ellipse((-180, 840, 620, 1540), fill=rgba("#D8EAFB", 68))
    od.rounded_rectangle((1400, 1060, 2330, 1240), radius=80, fill=rgba("#FFFFFF", 124))
    overlay = overlay.filter(ImageFilter.GaussianBlur(22))
    base.alpha_composite(overlay)
    draw = ImageDraw.Draw(base)

    draw.text((84, 76), REPO_NAME, font=load_font(76, bold=True), fill="#103A5D")
    intro_h = draw_wrapped(draw, (84, 176), INTRO, load_font(24), "#33475A", 670, 8)
    draw.text((84, 292 + intro_h), "GitHub", font=load_font(18, bold=True), fill="#8093A8")
    draw.rounded_rectangle((84, 322 + intro_h, 692, 370 + intro_h), radius=16, fill="#EAF3FF")
    draw.text((108, 334 + intro_h), REPO_URL, font=load_font(18), fill="#174B75")

    chip_x = 84
    chip_y = 396 + intro_h
    for label in ["综述重构", "图表说明", "公式解读", "模板保真", "导出核验"]:
        font = load_font(18, bold=True)
        tw = draw.textbbox((0, 0), label, font=font)[2]
        w = tw + 34
        draw.rounded_rectangle((chip_x, chip_y, chip_x + w, chip_y + 38), radius=16, fill="#F0F6FD")
        draw.text((chip_x + 17, chip_y + 9), label, font=font, fill="#133E66")
        chip_x += w + 12

    preview_img = Image.open(preview_path).convert("RGB")
    preview_panel(
        base,
        (790, 66, 2316, 786),
        "真实 PPT 效果预览",
        "展示完整页面缩略图，不裁切页面边界；海报中只保留代表性内容页，避免整套接触图带来的拥挤感。",
        preview_img,
    )

    cards = [
        (84, 514 + intro_h, 392, 674 + intro_h),
        (410, 514 + intro_h, 718, 674 + intro_h),
        (84, 692 + intro_h, 392, 852 + intro_h),
        (410, 692 + intro_h, 718, 852 + intro_h),
    ]
    for box, feature in zip(cards, FEATURES):
        feature_card(base, box, *feature)

    flow_panel(base, (84, 1078, 792, 1260))

    strip_img = Image.open(strip_path).convert("RGB")
    preview_panel(
        base,
        (820, 830, 1578, 1140),
        "局部页面放大",
        "三张页面均为完整缩略显示，用于突出结构页、建模页与控制页的层级与版式密度。",
        strip_img,
    )

    resource_panel(base, (1606, 830, 2316, 1260), FIGURES / "Fig_CPG_Hopf控制_深化版.png")

    out = POSTERS / "paper2scholarslides-poster-template-refined.png"
    base.convert("RGB").save(out, quality=96)
    return out


def add_textbox(slide, left, top, width, height, text, font_size, color, bold=False, name=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    if name:
        box.name = name
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*rgb(color))
    p.alignment = PP_ALIGN.LEFT
    tf.vertical_anchor = MSO_ANCHOR.TOP
    return box


def fill_shape(shape, fill_color, line_color=None, transparency=0.0, line_width=1.5):
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*rgb(fill_color))
    shape.fill.transparency = transparency
    if line_color:
        shape.line.color.rgb = RGBColor(*rgb(line_color))
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()


def place_picture_contain(slide, path: Path, left, top, width, height, padding=0.0):
    img = Image.open(path)
    box_w = float(width)
    box_h = float(height)
    pad = float(Inches(padding))
    inner_w = max(1.0, box_w - pad * 2)
    inner_h = max(1.0, box_h - pad * 2)
    src_ratio = img.width / img.height
    dst_ratio = inner_w / inner_h
    if src_ratio > dst_ratio:
        pic_w = inner_w
        pic_h = inner_w / src_ratio
    else:
        pic_h = inner_h
        pic_w = inner_h * src_ratio
    x = float(left) + (box_w - pic_w) / 2
    y = float(top) + (box_h - pic_h) / 2
    slide.shapes.add_picture(str(path), int(x), int(y), width=int(pic_w), height=int(pic_h))


def poster_pptx(preview_path: Path, strip_path: Path, png_path: Path) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bg = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    fill_shape(bg, "#F3F8FC")

    c1 = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(10.1), Inches(-0.55), Inches(4.4), Inches(4.4))
    fill_shape(c1, "#DCEBFB", transparency=0.42)
    c1.line.fill.background()
    c2 = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(-0.55), Inches(5.0), Inches(3.1), Inches(3.1))
    fill_shape(c2, "#E4F1FC", transparency=0.22)
    c2.line.fill.background()

    add_textbox(slide, Inches(0.48), Inches(0.36), Inches(4.2), Inches(0.8), REPO_NAME, 34, "#103A5D", True)
    add_textbox(slide, Inches(0.48), Inches(1.18), Inches(3.95), Inches(1.08), INTRO, 16, "#33475A")
    add_textbox(slide, Inches(0.48), Inches(2.36), Inches(0.8), Inches(0.22), "GitHub", 10, "#8093A8", True)
    url_box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.48), Inches(2.56), Inches(3.34), Inches(0.34))
    fill_shape(url_box, "#EAF3FF", None)
    add_textbox(slide, Inches(0.60), Inches(2.63), Inches(3.05), Inches(0.18), REPO_URL, 10, "#174B75")

    chip_x = 0.48
    for label in ["综述重构", "图表说明", "公式解读", "模板保真", "导出核验"]:
        width = 0.52 + len(label) * 0.13
        chip = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(chip_x), Inches(3.08), Inches(width), Inches(0.28))
        fill_shape(chip, "#F0F6FD", None)
        add_textbox(slide, Inches(chip_x + 0.08), Inches(3.15), Inches(width - 0.16), Inches(0.1), label, 9, "#133E66", True)
        chip_x += width + 0.07

    panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(4.35), Inches(0.34), Inches(8.35), Inches(4.02))
    fill_shape(panel, "#FFFFFF", "#D8E4F0", 0.0, 1.0)
    add_textbox(slide, Inches(4.58), Inches(0.52), Inches(2.8), Inches(0.3), "真实 PPT 效果预览", 18, "#123B61", True)
    add_textbox(slide, Inches(4.58), Inches(0.84), Inches(7.4), Inches(0.34), "展示完整页面缩略图，不裁切页面边界；海报中只保留代表性内容页，避免整套接触图带来的拥挤感。", 10, "#60768C")
    place_picture_contain(slide, preview_path, Inches(4.56), Inches(1.18), Inches(7.92), Inches(3.0), padding=0.06)

    card_positions = [
        (0.48, 3.76), (2.16, 3.76),
        (0.48, 5.06), (2.16, 5.06),
    ]
    for (x, y), (title, body, accent) in zip(card_positions, FEATURES):
        card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(1.56), Inches(1.15))
        fill_shape(card, "#FFFFFF", "#D8E4F0", 0.0, 1.0)
        dot = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x + 0.08), Inches(y + 0.08), Inches(0.18), Inches(0.18))
        fill_shape(dot, accent, None)
        add_textbox(slide, Inches(x + 0.32), Inches(y + 0.06), Inches(1.12), Inches(0.2), title, 12, "#123B61", True)
        add_textbox(slide, Inches(x + 0.08), Inches(y + 0.36), Inches(1.36), Inches(0.66), body, 9, "#314252")

    flow = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.48), Inches(6.02), Inches(4.28), Inches(1.0))
    fill_shape(flow, "#FFFFFF", "#D8E4F0", 0.0, 1.0)
    add_textbox(slide, Inches(0.66), Inches(6.20), Inches(1.0), Inches(0.2), "工作流程", 15, "#123B61", True)
    sx = 0.66
    for idx, label in enumerate(WORKFLOW):
        width = 0.64
        node = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(sx), Inches(6.57), Inches(width), Inches(0.32))
        fill_shape(node, "#FBFDFF", WORKFLOW_COLORS[idx], 0.0, 1.0)
        add_textbox(slide, Inches(sx), Inches(6.66), Inches(width), Inches(0.1), label, 8.5, WORKFLOW_COLORS[idx], True)
        if idx < len(WORKFLOW) - 1:
            x1 = Inches(sx + width)
            x2 = Inches(sx + width + 0.24)
            y = Inches(6.73)
            conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y, x2, y)
            conn.line.color.rgb = RGBColor(*rgb("#7CBAD2"))
            conn.line.width = Pt(1.5)
        sx += 0.95

    detail = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(4.5), Inches(4.50), Inches(3.35), Inches(2.57))
    fill_shape(detail, "#FFFFFF", "#D8E4F0", 0.0, 1.0)
    add_textbox(slide, Inches(4.68), Inches(4.66), Inches(1.6), Inches(0.2), "局部页面放大", 15, "#123B61", True)
    add_textbox(slide, Inches(4.68), Inches(4.92), Inches(2.9), Inches(0.3), "三张页面均为完整缩略显示，用于突出结构页、建模页与控制页的层级与版式密度。", 10, "#60768C")
    place_picture_contain(slide, strip_path, Inches(4.64), Inches(5.05), Inches(3.08), Inches(1.58), padding=0.03)

    res = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(8.0), Inches(4.50), Inches(4.7), Inches(2.57))
    fill_shape(res, "#FFFFFF", "#D8E4F0", 0.0, 1.0)
    add_textbox(slide, Inches(8.18), Inches(4.66), Inches(1.9), Inches(0.2), "图件资源与输出", 15, "#123B61", True)
    place_picture_contain(slide, FIGURES / "Fig_典型步态_中文线图.png", Inches(8.18), Inches(4.98), Inches(1.35), Inches(0.78), padding=0.02)
    place_picture_contain(slide, FIGURES / "Fig_CPG_Hopf控制_深化版.png", Inches(9.67), Inches(4.98), Inches(1.35), Inches(0.78), padding=0.02)
    place_picture_contain(slide, FIGURES / "Fig_LOS_MPC区别_深化版.png", Inches(11.16), Inches(4.98), Inches(1.35), Inches(0.78), padding=0.02)
    y_text = 5.66
    for line in ASSET_LINES:
        bullet = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(8.18), Inches(y_text + 0.02), Inches(0.08), Inches(0.08))
        fill_shape(bullet, "#2B6CB0", None)
        add_textbox(slide, Inches(8.34), Inches(y_text - 0.02), Inches(3.95), Inches(0.12), line, 9.5, "#314252")
        y_text += 0.23
    add_textbox(slide, Inches(11.18), Inches(6.36), Inches(1.15), Inches(0.5), "适用场景", 13, "#123B61", True)
    for idx, item in enumerate(USE_CASES):
        bullet = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(11.18), Inches(6.62 + idx * 0.14), Inches(0.07), Inches(0.07))
        fill_shape(bullet, "#1F8A70", None)
        add_textbox(slide, Inches(11.31), Inches(6.58 + idx * 0.14), Inches(1.2), Inches(0.1), item, 8.7, "#314252")

    out = POSTERS / "paper2scholarslides-poster-template-refined.pptx"
    prs.save(out)
    return out


def main() -> None:
    preview = clean_preview_grid()
    strip = detail_strip()
    poster = poster_png(preview, strip)
    pptx_path = poster_pptx(preview, strip, poster)
    print(preview)
    print(strip)
    print(poster)
    print(pptx_path)


if __name__ == "__main__":
    main()
