#!/usr/bin/env python
"""Scan a PPTX for process notes, filler phrasing, and task-specific banned terms."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from pptx import Presentation


DEFAULT_TERMS = [
    "user requested",
    "according to the user",
    "AI-generated",
    "paper library",
    "文献库",
    "用户要求",
    "根据用户",
    "面向用户",
    "制作说明",
    "AI生成",
    "链条",
    "痛点",
    "闭环",
    "任务评价",
    "虚无缥缈",
    "花里胡哨",
]


def iter_text(prs: Presentation):
    for slide_idx, slide in enumerate(prs.slides, 1):
        for shape_idx, shape in enumerate(slide.shapes, 1):
            if hasattr(shape, "text") and shape.text:
                text = shape.text.strip()
                if text:
                    yield slide_idx, shape_idx, text


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--term", action="append", default=[], help="Additional banned term")
    args = parser.parse_args()

    prs = Presentation(str(args.pptx))
    terms = DEFAULT_TERMS + args.term
    hits = []
    for slide_idx, shape_idx, text in iter_text(prs):
        normalized = text.replace("\n", " | ")
        for term in terms:
            if term and term in text:
                hits.append((slide_idx, shape_idx, term, normalized[:220]))

    print(f"slides={len(prs.slides)}")
    print(f"text_shapes={sum(1 for _ in iter_text(prs))}")
    print(f"hits={len(hits)}")
    for slide_idx, shape_idx, term, excerpt in hits:
        print(f"[slide {slide_idx:02d} shape {shape_idx:03d}] {term}: {excerpt}")
    return 1 if hits else 0


if __name__ == "__main__":
    sys.exit(main())
